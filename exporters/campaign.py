import json
from pathlib import Path
from typing import Any
from core.approval.gate import HumanApprovalGate


def export_approved_campaign(kit_id: str, kit: dict[str, Any], export_format: str, gate: HumanApprovalGate, output_dir: Path) -> dict[str, Any]:
    gate.require_approved(kit_id)
    fmt = export_format.lower().strip()
    output_dir.mkdir(parents=True, exist_ok=True)
    if fmt == "json":
        path = output_dir / f"{kit_id}.json"
        path.write_text(json.dumps(kit, indent=2, ensure_ascii=False), encoding="utf-8")
    elif fmt == "docx":
        from docx import Document
        path = output_dir / f"{kit_id}.docx"
        doc = Document()
        doc.add_heading(kit_id, level=1)
        doc.add_paragraph(json.dumps(kit, indent=2, ensure_ascii=False))
        doc.save(path)
    elif fmt == "xlsx":
        from openpyxl import Workbook
        path = output_dir / f"{kit_id}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Campaign Kit"
        ws.append(["Field", "Value"])
        for key, value in kit.items():
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)])
        wb.save(path)
    else:
        raise ValueError("Unsupported export_format. Use json, docx, or xlsx.")
    return {"status": "EXPORTED", "format": fmt, "path": str(path), "approved": True}
